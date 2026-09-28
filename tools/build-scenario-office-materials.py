#!/usr/bin/env python3
"""Generate scenario originals from the canonical Markdown without changing terms.

Run with a Python environment containing python-docx, reportlab and pypdf.
Pass a TTF font that covers Chinese with --font. Rendering is a separate QA gate.
"""

import argparse
import hashlib
import json
import re
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


ROOT = Path(__file__).resolve().parents[1]
MATERIALS = ROOT / "scenarios/sales-contract-handoff/materials"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def paragraphs(path):
    result = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        kind = "body"
        if line.startswith("### "):
            kind, line = "h2", line[4:]
        elif line.startswith("## "):
            kind, line = "h1", line[3:]
        elif line.startswith("# "):
            kind, line = "title", line[2:]
        elif line.startswith("> "):
            kind, line = "notice", line[2:]
        elif line.startswith("- "):
            line = line[2:]
        result.append((kind, line))
    return result


def plain(text):
    return text.replace("**", "")


def normalized(text):
    return re.sub(r"\s+", "", text)


def assert_content(path, expected, actual):
    if normalized(expected) != normalized(actual):
        raise ValueError(f"Text differs from canonical Markdown: {path.name}")


def make_docx(source, output):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(18)
    section.left_margin = section.right_margin = Mm(22)
    for name, size in (("Normal", 10.5), ("Title", 19), ("Heading 1", 12), ("Heading 2", 10.5)):
        style = document.styles[name]
        style.font.name = "Arial Unicode MS"
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for key in list(fonts.attrib):
            if "theme" in key.lower():
                del fonts.attrib[key]
        fonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        fonts.set(qn("w:cs"), "Arial Unicode MS")
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.25
    document.styles["Title"].paragraph_format.space_after = Pt(12)
    document.styles["Heading 1"].paragraph_format.space_before = Pt(9)
    document.styles["Heading 2"].paragraph_format.space_before = Pt(5)
    document.core_properties.author = "MVP1 场景测试"
    document.core_properties.title = plain(source[0][1])
    for kind, text in source:
        style = {"title": "Title", "h1": "Heading 1", "h2": "Heading 2"}.get(kind, "Normal")
        paragraph = document.add_paragraph(style=style)
        paragraph.paragraph_format.keep_together = True
        for index, part in enumerate(re.split(r"\*\*", text)):
            run = paragraph.add_run(part)
            run.font.name = "Arial Unicode MS"
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial Unicode MS")
            if index % 2:
                run.bold = True
            if kind == "notice":
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(75, 75, 75)
    document.save(output)
    expected = "\n".join(plain(text) for _, text in source)
    assert_content(output, expected, "\n".join(p.text for p in Document(output).paragraphs))


def make_pdf(source, output, font):
    pdfmetrics.registerFont(TTFont("ScenarioCJK", str(font)))
    pdfmetrics.registerFontFamily("ScenarioCJK", normal="ScenarioCJK", bold="ScenarioCJK")
    body = ParagraphStyle("body", fontName="ScenarioCJK", fontSize=10.5, leading=16,
                          spaceAfter=7, wordWrap="CJK")
    styles = {
        "body": body,
        "notice": ParagraphStyle("notice", parent=body, fontSize=9, leading=14, textColor=colors.HexColor("#4b4b4b")),
        "title": ParagraphStyle("title", parent=body, fontSize=19, leading=25, spaceAfter=14, keepWithNext=True),
        "h1": ParagraphStyle("h1", parent=body, fontSize=12, leading=18, spaceBefore=12, spaceAfter=5, keepWithNext=True),
        "h2": ParagraphStyle("h2", parent=body, fontSize=11, leading=17, spaceBefore=7, spaceAfter=4, keepWithNext=True),
    }
    story = []
    for kind, text in source:
        rich = "".join((f"<b>{escape(part)}</b>" if i % 2 else escape(part))
                       for i, part in enumerate(re.split(r"\*\*", text)))
        story.append(Paragraph(rich, styles[kind]))
    story.append(Spacer(1, 3))
    SimpleDocTemplate(str(output), pagesize=A4, leftMargin=62, rightMargin=62,
                      topMargin=51, bottomMargin=51, title=plain(source[0][1]),
                      author="MVP1 场景测试").build(story)
    reader = PdfReader(output)
    assert_content(output, "\n".join(plain(t) for _, t in source),
                   "\n".join(page.extract_text() for page in reader.pages))
    return len(reader.pages)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--font", type=Path, required=True)
    args = parser.parse_args()
    if not args.font.is_file():
        parser.error("--font must name a readable Chinese-capable TTF file")
    sources = [MATERIALS / "合同样例.md", MATERIALS / "技术协议样例.md"]
    outputs = [MATERIALS / "合同样例.docx", MATERIALS / "技术协议样例.pdf"]
    make_docx(paragraphs(sources[0]), outputs[0])
    pdf_pages = make_pdf(paragraphs(sources[1]), outputs[1], args.font)
    manifest = {
        "schemaVersion": 1,
        "purpose": "124 场景流程测试，不构成真实合同、技术承诺或签署凭证",
        "status": "prepared_not_dispatched",
        "fontSha256": digest(args.font),
        "materials": [
            {"source": src.name, "sourceSha256": digest(src), "file": out.name,
             "mediaType": mime, "bytes": out.stat().st_size, "sha256": digest(out),
             "textMatchesSource": True}
            for src, out, mime in zip(sources, outputs, (
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/pdf"))
        ],
        "expectedFindings": ["交付期限：合同 20 个日历日，技术协议 15 个日历日",
                             "连续运行验收：合同 24 小时，技术协议 72 小时"],
        "pdfPages": pdf_pages,
    }
    for entry in manifest["materials"]:
        if entry["bytes"] > 2 * 1024 * 1024:
            raise ValueError(f"Original exceeds the 2 MiB contract limit: {entry['file']}")
    (MATERIALS / "办公材料清单.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
